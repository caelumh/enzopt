"""Build this project's data/ from the plate-reader workbook of Wang et al. (2025) (figshare
doi:10.6084/m9.figshare.26961733, CC BY 4.0; see README.md).

    uv run --with openpyxl python -I projects/ldh_wang2025/make_data.py

raw/Experimental_data.xlsx holds five kinetic runs of rabbit-muscle LDH (pyruvate + NADH -> lactate + NAD+, PBS,
23 °C), each a sheet "Kinetik <k>. <rows>" with the raw A340 of the whole 96-well plate every 10 s for 490 s, and an
analysis sheet "Auswertung Kinetik <k>." with the conditions: NADH, enzyme (mg/mL), the highest pyruvate (2 mM) and the
dilution factor (1.5). In each run the reaction rows (D-F or A-C, triplicates) hold pyruvate 2 mM / 1.5^(column - 1)
in columns 1-12; rows G-H are empty wells (plate background).

Each reaction well becomes data/r<run>_c<column>_<row>.csv (time_s, NADH_mM): NADH = (A340 - background) / (ε d) with
the authors' ε = 6.3 /(mM cm) and d = 0.3118 cm (100 µL in a 6.39 mm well, analysis sheet); background = mean of the
empty wells G-H at the first reading. The first reading is the initial NADH (it is about 20 % below the nominal NADH;
README); pyruvate is the nominal value, lactate and NAD+ zero. The enzyme is converted to active sites with 36.6 kDa
per subunit (as in ldh_holmes1988). Run 3's enzyme is written as 0.125 mg/mL in its analysis sheet; it is read as
0.000125 mg/mL (the other runs vary one factor around 0.000125 mg/mL, and run 3 varies only NADH).
"""

from __future__ import annotations

import csv
import sys
from pathlib import Path

import numpy as np
import openpyxl

HERE = Path(__file__).resolve().parent
RAW = HERE / "raw" / "Experimental_data.xlsx"
EPS_D = 6.3 * 0.3118  # A340 per mM NADH (authors' epsilon and path length)
KDA_PER_SITE = 36.6
ENZYME_FIX = {3: 0.000125}  # mg/mL, run 3's sheet says 0.125
ROLE = {1: "holdout", 2: "train", 3: "train", 4: "train", 5: "holdout"}  # holdout: lower and higher enzyme


def well_rows(ws) -> tuple[np.ndarray, dict[str, np.ndarray]]:
    rows = list(ws.iter_rows(values_only=True))
    t = np.array([x for x in next(r for r in rows if r[0] == "Time")[1:] if isinstance(x, (int, float))], float)
    wells = {}
    for r in rows:
        name = r[0]
        if isinstance(name, str) and len(name) <= 3 and name[:1] in "ABCDEFGH" and name[1:].isdigit():
            wells[name] = np.array([x for x in r[1:] if isinstance(x, (int, float))], float)
    return t, wells


def conditions(ws) -> dict[str, float]:
    out = {}
    for r in ws.iter_rows(min_row=25, max_row=45, values_only=True):
        if isinstance(r[2], str) and isinstance(r[4], (int, float)):
            key = r[2].split()[0].lower()  # "c", "enzyme", "dilution", "highest"
            out[key] = float(r[4])
    return out


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    out = HERE / "data"
    out.mkdir(exist_ok=True)
    wb = openpyxl.load_workbook(RAW, read_only=True, data_only=True)
    lines, noise = [], []
    for k in range(1, 6):
        ws = next(w for w in wb.worksheets if w.title.startswith(f"Kinetik {k}."))
        cond = conditions(wb[f"Auswertung Kinetik {k}."])
        nadh_nom, top, dil = cond["c"], cond["highest"], cond["dilution"]
        mg_ml = ENZYME_FIX.get(k, cond["enzyme"])
        et_nm = mg_ml / KDA_PER_SITE * 1e6
        rows = ws.title.split()[-1]
        letters = [chr(c) for c in range(ord(rows[0]), ord(rows[-1]) + 1)]
        t, wells = well_rows(ws)
        background = float(np.mean([wells[f"{r}{c}"][0] for r in "GH" for c in range(1, 13)]))
        lines.append(
            f"  # run {k} (sheet {ws.title!r}): NADH {nadh_nom:g} mM nominal, enzyme {mg_ml:g} mg/mL = {et_nm:.3g} nM "
            f"active sites, background A340 {background:.3f}"
        )
        for c in range(1, 13):
            pyr = top / dil ** (c - 1)
            for r in letters:
                a = wells.get(f"{r}{c}")
                if a is None or len(a) != len(t):
                    continue
                y = (a - background) / EPS_D
                name = f"r{k}_c{c:02d}_{r}"
                with (out / f"{name}.csv").open("w", newline="", encoding="utf-8") as fh:
                    w = csv.writer(fh, lineterminator="\n")
                    w.writerow(["time_s", "NADH_mM"])
                    w.writerows([f"{ti:g}", f"{yi:.6f}"] for ti, yi in zip(t, y, strict=True))
                noise.append(float(np.std(np.diff(y, 2)) / np.sqrt(6)))
                lines.append(
                    f"  {name}: {{file: data/{name}.csv, NADH: {y[0]:.5g}, Pyr: {pyr:.5g}, Lac: 0, NAD: 0, "
                    f"Et_nM: {et_nm:.4g}, role: {ROLE[k]}}}"
                )
    print("\n".join(lines))
    q = np.percentile(noise, [10, 50, 90]) * 1e3
    spread = "/".join(f"{x:.2f}" for x in q)
    print(f"# {len(noise)} wells; read noise (sd of second differences / sqrt 6) p10/p50/p90: {spread} µM")


if __name__ == "__main__":
    main()
